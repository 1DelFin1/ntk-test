import io
import re
import email
from email import policy
from html import unescape

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook


def parse_txt(data: bytes) -> str:
    text = data.decode("utf-8-sig", errors="replace")
    if "\uFFFD" in text:
        try:
            text = data.decode("cp1251")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
    return text


def parse_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise RuntimeError("PDF защищён паролем.")
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def parse_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]

    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))

    return "\n".join(parts)


def parse_xlsx(data: bytes) -> str:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    parts = []

    for ws in wb.worksheets:
        parts.append(f"--- Лист: {ws.title} ---")
        for row in ws.iter_rows(values_only=True):
            if any(c is not None for c in row):
                parts.append(" | ".join("" if c is None else str(c) for c in row))

    return "\n".join(parts)


_HTML_TAG_RE = re.compile(r"<[^>]+>")


def _strip_html(html: str) -> str:
    html = html.replace("<br>", "\n").replace("<br/>", "\n").replace("<br />", "\n")
    return unescape(_HTML_TAG_RE.sub("", html)).strip()


def parse_eml(data: bytes) -> str:
    msg = email.message_from_bytes(data, policy=policy.default)
    parts = []

    if msg["subject"]:
        parts.append(f"Тема: {msg['subject']}")
    if msg["from"]:
        parts.append(f"От: {msg['from']}")
    if msg["to"]:
        parts.append(f"Кому: {msg['to']}")

    if msg.is_multipart():
        text_parts = []
        html_parts = []
        for part in msg.walk():
            ctype = part.get_content_type()
            payload = part.get_payload(decode=True)
            if not payload:
                continue
            charset = part.get_content_charset() or "utf-8"
            try:
                decoded = payload.decode(charset, errors="replace")
            except (LookupError, ValueError):
                decoded = payload.decode("utf-8", errors="replace")
            if ctype == "text/plain":
                text_parts.append(decoded)
            elif ctype == "text/html":
                html_parts.append(_strip_html(decoded))
        body = "\n".join(text_parts) if text_parts else "\n".join(html_parts)
    else:
        raw = msg.get_content() or ""
        body = _strip_html(raw) if msg.get_content_type() == "text/html" else raw

    if body.strip():
        parts.append(body)

    return "\n".join(parts)


def extract_text_from_file(filename: str, data: bytes) -> str:
    name = filename.lower()

    if name.endswith(".pdf"):
        return parse_pdf(data)
    if name.endswith(".docx"):
        return parse_docx(data)
    if name.endswith((".xlsx", ".xlsm")):
        return parse_xlsx(data)
    if name.endswith(".eml"):
        return parse_eml(data)

    return parse_txt(data)